# -*- coding: utf-8 -*-
"""Build the Scintilla GT3 pack author-feedback Word doc."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(
    r"C:\Users\anton\AppData\Local\BeamNG\BeamNG.drive\current\mods\unpacked"
    r"\scintilla_gt3\AUTHOR_FEEDBACK_local_fixes.docx"
)


def set_run(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, size=16 if level == 1 else 13, bold=True, color=(0x1F, 0x3A, 0x5F))
    return p


def add_p(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, bold=bold)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run(run)
        p.paragraph_format.space_after = Pt(3)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    t = doc.add_paragraph()
    r = t.add_run("Scintilla GT3 Racing Parts — tester notes for the authors")
    set_run(r, size=20, bold=True, color=(0x1F, 0x3A, 0x5F))
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT

    sub = doc.add_paragraph()
    r = sub.add_run(
        "Local unpacked copy only. Not an official pack release.\n"
        "Pack: Scintilla GT3 (Exchy / Turbo49 / Cyborella / ©yborella / _N_S_)\n"
        "Game: BeamNG.drive 0.39.x  |  Date: 19 August 2026\n"
        "From: local tester (NTF tire work on this car)"
    )
    set_run(r, size=11, color=(0x44, 0x44, 0x44))

    add_heading(doc, "Purpose", 1)
    add_p(
        doc,
        "This note is a growing list of pack issues found while using the local GT3 copy "
        "for tire-physics testing. Fixes described here are already applied in the unpacked "
        "local files so testing can continue. They are offered as suggested upstream changes, "
        "not as a demand to take the local tree as-is.",
    )

    add_heading(doc, "1. Traction control lighting ABS and locking fronts (no brake pedal)", 1)
    add_heading(doc, "What we saw", 2)
    add_p(
        doc,
        "On long stints (Belasco), the front-left tire would flatspot (about 12–17%) with the "
        "brake pedal not applied. Pitwall showed a lock / hot-flash on that corner. It looked "
        "like TC during turn-in was firing ABS on the fronts.",
    )
    add_heading(doc, "Is this based on a real GT3 system?", 2)
    add_p(
        doc,
        "Partly. Real FIA GT3 customer cars have driver-adjustable ABS and traction control. "
        "FIA Appendix J, Article 257A, Art. 505 (Driving aids) is explicit:",
    )
    add_p(
        doc,
        "“Any electronic stability control system is forbidden. Only traction control systems "
        "managing the engine output are permitted. The control of other devices is forbidden "
        "unless specifically homologated.”",
        bold=True,
    )
    add_p(
        doc,
        "So: engine-torque TC + homologated ABS = race-correct. Brake-based ESC / ESP that "
        "stabs an undriven front to catch yaw = road-car DSC, and it is what caused the lockups. "
        "BeamNG stock DSE is a road-car ESP stack (yawControl + brakeControl). Wiring that "
        "straight into a GT3 electronics slot is the mismatch, not a bad idea for rotary TC/ABS.",
    )
    add_heading(doc, "Root cause in the local pack", 2)
    add_bullets(
        doc,
        [
            "File: vehicles/scintilla/mechanical/scintilla_SSC_gt3.jbeam",
            "Drive mode “TC & ESC On” enabled yawControl on brakeControl. DSE then applies brake torque to individual corners (including fronts) with pedal = 0.",
            "ABS is on every pressure wheel (enableABS, default level 4 → slipRatioTarget 0.25). An ESC stab exceeds that target → ABS light, lock, flatspot.",
            "In the same corner, RWD TC also lights (inside-rear spin). The two lamps fire together, so it feels like “TC activated ABS.” It is ESC brake + ABS.",
            "The TC part also had brakeControl.useForTractionControl = true at spawn (placeholder slip 99). Mode On tried to turn that off; if the mode apply is late, TC can brake too. Front locks still point at yaw, not TC.",
            "Legacy lua/controller/gt3TractionControl.lua is unused (zip leftover). Current electronics are stock DSE.",
        ],
    )
    add_heading(doc, "Options considered", 2)
    add_p(doc, "A. Correct stock DSE to FIA GT3 (chosen — applied locally)", bold=True)
    add_p(
        doc,
        "Keep BeamNG tractionControl + motorTorqueControl (engine cut) and ABS. Disable "
        "yawControl on brakes, motors, AWD, diff, and aero. Rename the drive mode to “TC On / TC Off.” "
        "Keep the author’s TC 1–10 and ABS 1–7 rotaries.",
    )
    add_p(doc, "Pros: matches Art. 505; one-file jbeam change; no new Lua; author’s UI stays; stops uncommanded front locks.")
    add_p(
        doc,
        "Cons: the car will rotate more (no ESP catch); ESC Tuning sliders no longer do anything; "
        "less “safe” for casual players than road DSC.",
    )
    add_p(doc, "B. Write a custom GT3 TC (axle-speed Lua, like the unused gt3TractionControl.lua)")
    add_p(
        doc,
        "Pros: full control of cut shape. Cons: axle-average TC false-triggers in corners "
        "(inside/outside speed split); more code to maintain vs 0.39 DSE; worse than stock "
        "motorTorqueControl slip on the driven axle. Not recommended.",
    )
    add_p(doc, "C. Keep brake-yaw ESC but raise thresholds")
    add_p(
        doc,
        "Pros: still “catches” spins. Cons: still not GT3-legal; still fights ABS; would need "
        "endless PID tuning. Rejected.",
    )
    add_heading(doc, "Local fix applied (please consider upstream)", 2)
    add_bullets(
        doc,
        [
            "brakeControl.useForTractionControl = false and useForYawControl = false (TC never brakes; ESC never stabs wheels).",
            "motorTorqueControl.useForYawControl = false; TC remains torque-only on mainEngine.",
            "Drive mode On: yawControl.isEnabled = false on yaw, brake, motor, AWD, diff, aero. TC stays on.",
            "Drive mode names: “TC On” / “TC Off”. ESC dash lamp stays dim.",
            "Respawn after the change. Confirm: corner on throttle, pedal up, fronts must not lock or flatspot. TC lamp may still flash on inside-rear spin — that is throttle cut only.",
        ],
    )

    add_heading(doc, "2. Left/right pad weight (driver vs ballast)", 1)
    add_p(
        doc,
        "The left-seat driver mass sits on +X (FL). Without a matching right-side mass the static "
        "pad is left-heavy. Raising a spring perch only unloads that corner; it does not cancel "
        "driver mass. The correct GT3 BoP tool is passenger-side ballast on chassis node f2r "
        "(negative X, right floor): nodeWeight = $ballast + 4.5.",
    )
    add_bullets(
        doc,
        [
            "Local chassis default $ballast is 10 kg — far too light vs a 70–80 kg driver. Anyone who spawns without a .pc override will sit left-heavy.",
            "gt3_balanced.pc and gt3_endurance.pc set $weight_driver 80 and $ballast 95 so pad L/R can even. That is the intended method.",
            "Spring perch variables were documented locally so testers stop chasing L/R with ride height.",
            "Suggest: raise the jbeam default $ballast nearer to default driver weight, or default the ballast slider to the same number as $weight_driver in the configs.",
        ],
    )

    add_heading(doc, "3. BeamNG 0.39 local hardening", 1)
    add_p(
        doc,
        "The pack was written against older controller load order. 0.39 is stricter about when "
        "peer controllers exist. Local hardening (already in the unpacked copy):",
    )
    add_bullets(
        doc,
        [
            "lua/controller/publicMethods_gt3.lua — do not bind gt3Timer / gt3Gauges / gt3PitLimiter at require-time. Bind in init / initLastStage via getControllerSafe. Prevents nil peer crashes when input actions fire before all controllers exist.",
            "lua/controller/gt3PitLimiter.lua — sendCurrentState renamed to ASCII (local 0.39). Non-ASCII identifiers in that path were a load risk.",
            "scintilla_chassis_gt3.jbeam — slot defaults pointed at GT3 body/hood/fenders/undertray/tanks so balanced/endurance configs actually get the GT3 aero homologation parts on 0.39 spawn instead of falling through to street pieces.",
        ],
    )

    add_heading(doc, "4. Leftovers / small notes", 1)
    add_bullets(
        doc,
        [
            "gt3TractionControl.lua is still listed in mod_info but is not slotted. Safe to omit from a future zip, or keep as unused archive.",
            "ABS remains adjustable (1–7) and should stay — real GT3 ABS maps exist. The bug was ESC using the same brakes ABS is watching.",
            "Tire flatspots during NTF testing on this car were electronics, not the tire mod. After the DSE fix, a 22 km stint should not pick up FL lockups with the pedal up.",
        ],
    )

    add_heading(doc, "Suggested test after merge", 1)
    add_bullets(
        doc,
        [
            "Spawn GT3 Endurance or Balanced, DSE = TC On, ABS mid (4).",
            "Belasco or any tight right/left: throttle on, brake 0. Front wheelspeed must stay with airspeed. No FL/FR flatspot channel.",
            "Hard launch: TC should cut engine (rears may spin briefly). No front ABS chatter.",
            "Hard braking in a straight: ABS should pulse as before.",
            "DSE = TC Off: full wheelspin available; ABS still on if the ABS part is fitted.",
        ],
    )

    add_heading(doc, "Files touched locally", 1)
    add_bullets(
        doc,
        [
            "vehicles/scintilla/mechanical/scintilla_SSC_gt3.jbeam (this TC/ESC fix)",
            "vehicles/scintilla/bodywork/scintilla_chassis_gt3.jbeam (ballast copy + 0.39 slot defaults)",
            "vehicles/scintilla/mechanical/scintilla_suspension_F_gt3.jbeam / _R_gt3.jbeam (perch vs ballast copy)",
            "vehicles/scintilla/lua/controller/publicMethods_gt3.lua",
            "vehicles/scintilla/lua/controller/gt3PitLimiter.lua",
            "gt3_balanced.pc / gt3_endurance.pc ($ballast 95 / $weight_driver 80)",
        ],
    )

    footer = doc.add_paragraph()
    r = footer.add_run(
        "Happy to walk through a diff or a short video of the pre-fix FL lock vs post-fix corner. "
        "Thank you for the pack — it is the car we use for high-aero tire work."
    )
    set_run(r, size=11, color=(0x44, 0x44, 0x44))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
